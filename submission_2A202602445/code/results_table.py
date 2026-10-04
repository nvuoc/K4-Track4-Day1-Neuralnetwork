"""results_table.py — Lưu kết quả và xuất bảng experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import glob
import json
import os
from pathlib import Path

import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có.
    """
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_path = os.path.join(results_dir, f"{exp_id}.json")

    payload = {
        "cfg": result.get("cfg", {}),
        "history": result.get("history", {}),
        "summary": result.get("summary", {}),
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)

    return out_path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    pattern = os.path.join(results_dir, "*.json")
    files = sorted(glob.glob(pattern))
    results = []
    for fpath in files:
        with open(fpath, "r", encoding="utf-8") as f:
            data = json.load(f)
            results.append(data)
    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.
    """
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")

    hidden = cfg.get("hidden", (256, 128))
    if isinstance(hidden, (list, tuple)):
        hidden_str = "-".join(str(h) for h in hidden)
    else:
        hidden_str = str(hidden)

    clip_norm = cfg.get("clip_norm", None)
    clip_str = "none" if clip_norm is None else clip_norm

    eval_acc = None
    eval_f1 = None
    if eval_scores is not None:
        eval_acc = eval_scores.get("accuracy", None)
        eval_f1 = eval_scores.get("macro_f1", None)

    return {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce").upper(),
        "optimizer": cfg.get("optimizer", "sgd_momentum"),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": clip_str,
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", ""),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_acc if eval_acc is not None else "",
        "eval_macro_f1": eval_f1 if eval_f1 is not None else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes if notes else cfg.get("notes", ""),
    }


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước:
      1. wb = openpyxl.load_workbook(template_path)  (KHÔNG dùng data_only=True)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức
      4. wb.save(out_path)
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    formula_cols = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

    # Đọc header dòng 1
    col_mapping = {}
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(1, col_idx).value
        if val is not None:
            col_mapping[str(val).strip()] = col_idx

    # Ghi dữ liệu từ dòng 2
    for i, row in enumerate(rows):
        row_num = 2 + i
        for key, value in row.items():
            if key in formula_cols:
                continue
            if key in col_mapping:
                col_idx = col_mapping[key]
                ws.cell(row=row_num, column=col_idx, value=value)

    # Đảm bảo thư mục lưu tồn tại
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    print(f"Đã lưu bảng kết quả: {out_path} ({len(rows)} thí nghiệm)")
